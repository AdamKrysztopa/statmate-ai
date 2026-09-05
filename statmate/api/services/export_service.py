"""Utilities for exporting analysis results to PDF, DOCX, and CSV."""

from __future__ import annotations

import base64
import binascii
import csv
import html
import io
import json
import zipfile
from datetime import datetime
from typing import Any

import pandas as pd

from statmate.core.pii import mask_dataframe


class ExportService:
    """Build lightweight exports for analyses."""

    @staticmethod
    def _html_safe(text: str | None) -> str:
        """Escape text for use in HTML text *and* attribute contexts.

        Quotes are escaped as well as angle brackets: report content (dataset names,
        LLM-authored plot titles, column names) is interpolated into quoted attributes,
        where an unescaped quote would allow attribute/CSS injection into the document
        handed to WeasyPrint.
        """
        return html.escape(str(text) if text is not None else '', quote=True)

    #: MIME types permitted in report image data URIs.
    _ALLOWED_IMAGE_TYPES = frozenset({'image/png', 'image/jpeg', 'image/svg+xml'})

    @classmethod
    def _image_data_uri(cls, content_type: str | None, image_base64: str | None) -> str | None:
        """Build a validated ``data:`` URI for a report image.

        HTML escaping stops markup breakout but does not make a URI well-formed: the MIME
        type and payload still reach WeasyPrint's URL handling. This whitelists the MIME
        type and round-trips the payload through a strict base64 decode/encode, so anything
        malformed or smuggling extra URI syntax is dropped rather than rendered.

        Args:
            content_type: Declared MIME type of the image.
            image_base64: Base64-encoded image payload.

        Returns:
            A safe ``data:`` URI, or None when the input is missing or invalid.
        """
        if not image_base64:
            return None
        mime = (content_type or 'image/png').strip().lower()
        if mime not in cls._ALLOWED_IMAGE_TYPES:
            mime = 'image/png'
        try:
            raw = base64.b64decode(str(image_base64), validate=True)
        except (ValueError, binascii.Error):
            return None
        if not raw:
            return None
        return f'data:{mime};base64,{base64.b64encode(raw).decode("ascii")}'

    @classmethod
    def _plot_src(cls, plot: dict[str, Any]) -> str:
        """Return a validated data URI for a plot, or an empty string if unusable."""
        return cls._image_data_uri(plot.get('content_type'), plot.get('image_base64')) or ''

    @staticmethod
    def _blocked_url_fetcher(url: str) -> dict[str, Any]:
        """Reject every non-``data:`` URL WeasyPrint tries to fetch.

        Report HTML embeds all of its images inline, so the renderer never has a
        legitimate reason to reach the network. Refusing outbound fetches closes the
        SSRF and CSS-injection paths in WeasyPrint that have no upstream fix.

        Args:
            url: URL WeasyPrint asked to resolve.

        Returns:
            The fetch result for permitted ``data:`` URLs.

        Raises:
            ValueError: If the URL is not a ``data:`` URL.
        """
        if not url.lower().startswith('data:'):
            raise ValueError(f'Blocked non-data URL in report rendering: {url[:60]}')
        from weasyprint.urls import default_url_fetcher

        return default_url_fetcher(url)

    @staticmethod
    def _latex_escape(text: str | None) -> str:
        """Escape LaTeX-reserved characters to avoid compilation failures."""
        if text is None:
            return ''
        replacements = {
            '\\': r'\textbackslash{}',
            '&': r'\&',
            '%': r'\%',
            '$': r'\$',
            '#': r'\#',
            '_': r'\_',
            '{': r'\{',
            '}': r'\}',
            '~': r'\textasciitilde{}',
            '^': r'\textasciicircum{}',
        }
        escaped = str(text)
        for target, replacement in replacements.items():
            escaped = escaped.replace(target, replacement)
        return escaped

    @classmethod
    def render_html_report(cls, analysis: dict[str, Any]) -> str:
        """Render an HTML report that can be used for PDF/DOCX conversion."""
        summary = cls._html_safe(analysis.get('summary') or analysis.get('results_detail', {}).get('summary'))
        probabilities = analysis.get('probabilities') or {}
        effect_sizes = analysis.get('effect_sizes') or analysis.get('results_detail', {}).get('effect_sizes') or {}
        plots = analysis.get('plots') or analysis.get('results_detail', {}).get('plots') or []
        decision_steps = (
            analysis.get('decision_steps') or analysis.get('results_detail', {}).get('decision_steps') or []
        )
        workflow_graph = (
            analysis.get('workflow_graph') or analysis.get('results_detail', {}).get('workflow_graph') or {}
        )
        graph_assets = workflow_graph.get('assets') or {}

        table_rows = ''.join(
            f"<tr><td>{cls._html_safe(name)}</td>"
            f"<td>{probabilities[name]:.4f}</td>"
            f"<td>{'Significant' if probabilities[name] < 0.05 else 'Not significant'}</td></tr>"
            for name in probabilities
        )

        effect_rows = ''.join(
            f'<tr><td>{cls._html_safe(name)}</td><td>{value:.3f}</td></tr>' for name, value in effect_sizes.items()
        )

        plot_blocks = ''.join(
            f"""
            <div class="plot">
              <div class="plot-title">{cls._html_safe(plot.get("title") or "Plot")}</div>
              <img src="{cls._html_safe(cls._plot_src(plot))}" />
              <div class="plot-meta">{cls._html_safe(plot.get("description") or "")}</div>
            </div>
            """
            for plot in plots
        )

        steps = ''.join(
            f"<li><strong>{cls._html_safe(step.get('step'))}</strong> — {cls._html_safe(step.get('detail'))}</li>"
            for step in decision_steps
        )

        step_table_rows = ''.join(
            f"<tr>"
            f"<td>{cls._html_safe(step.get('step'))}</td>"
            f"<td>{cls._html_safe(step.get('detail'))}</td>"
            f"<td>{cls._html_safe(step.get('timestamp'))}</td>"
            "<td>" + ''
            if step.get('p_value') is None
            else f"{step.get('p_value'):.4g}" + "</td>"
            f"<td>{cls._html_safe(str(step.get('progress_pct'))) if step.get('progress_pct') is not None else ''}</td>"
            f"</tr>"
            for step in decision_steps
        )

        generated_at = datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')
        dataset_name = analysis.get('dataset_name') or 'Dataset'
        header_title = cls._html_safe(dataset_name)

        return f"""
        <html>
          <head>
            <style>
              @page {{ size: A4; margin: 24px; }}
              body {{ font-family: Arial, sans-serif; color: #0f172a; padding: 24px; }}
              h1 {{ margin-bottom: 6px; }}
              .muted {{ color: #475569; }}
              img {{ max-width: 100%; height: auto; }}
              table {{ width: 100%; border-collapse: collapse; margin: 12px 0; }}
              th, td {{ border: 1px solid #e2e8f0; padding: 8px; text-align: left; }}
              th {{ background: #f8fafc; }}
              .plot-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(240px, 1fr)); gap: 12px; }}
              .plot {{ border: 1px solid #e2e8f0; border-radius: 10px; overflow: hidden;
              page-break-inside: avoid; break-inside: avoid; }}
              .plot img {{ width: 100%; height: auto; display: block; page-break-inside: avoid; break-inside: avoid; }}
              .plot-title {{ font-weight: 700; padding: 8px; background: #f8fafc; }}
              .plot-meta {{ padding: 8px; color: #475569; }}
              .graph-card {{ border: 1px solid #e2e8f0; border-radius: 12px; padding: 12px;
              margin: 8px 0 16px; background: #f8fafc; }}
              .graph-card img {{ width: 100%; max-height: 480px; object-fit: contain; }}
            </style>
          </head>
          <body>
            <h1>Analysis Report · {header_title}</h1>
            <p class="muted">Generated at {generated_at}</p>
            <h2>Summary</h2>
            <p>{summary or "No summary available."}</p>
            <h2>Workflow graph</h2>
            <div class="graph-card">
              {
            "<img alt='"
            + cls._html_safe(graph_assets.get("alt") or "Workflow graph")
            + "' src='"
            + cls._html_safe(
                cls._image_data_uri("image/png", graph_assets.get("png_base64"))
                or cls._image_data_uri("image/svg+xml", graph_assets.get("svg_base64"))
                or ""
            )
            + "' />"
            if graph_assets.get("svg_base64")
            else "<div class='muted'>No workflow graph available.</div>"
        }
            </div>
            <h2>Statistical tests</h2>
            <table>
              <thead><tr><th>Test</th><th>p-value</th><th>Callout</th></tr></thead>
              <tbody>
                {table_rows or '<tr><td colspan="3">No tests reported.</td></tr>'}
              </tbody>
            </table>
            <h2>Effect sizes</h2>
            <table>
              <thead><tr><th>Metric</th><th>Value</th></tr></thead>
              <tbody>
                {effect_rows or '<tr><td colspan="2">No effect sizes computed.</td></tr>'}
              </tbody>
            </table>
            <h2>Plots</h2>
            <div class="plot-grid">
              {plot_blocks or '<div class="muted">No plots generated.</div>'}
            </div>
            <h2>Decision path</h2>
            <ol>{steps or "<li>No steps captured.</li>"}</ol>
            <h2>Execution steps</h2>
            <table>
              <thead><tr><th>Step</th><th>Detail</th><th>Timestamp</th><th>p-value</th><th>Progress %</th></tr></thead>
              <tbody>
                {step_table_rows or '<tr><td colspan="5">No steps captured.</td></tr>'}
              </tbody>
            </table>
          </body>
        </html>
        """

    @staticmethod
    def render_pdf(html: str) -> bytes:
        """Render HTML to PDF using WeasyPrint."""
        try:
            from weasyprint import HTML
        except Exception as exc:  # pragma: no cover - optional dependency
            raise RuntimeError('WeasyPrint is required for PDF export') from exc
        return HTML(string=html, url_fetcher=ExportService._blocked_url_fetcher).write_pdf()

    @classmethod
    def render_latex_report(cls, analysis: dict[str, Any]) -> bytes:
        """Render a minimal LaTeX report for offline use."""
        summary = cls._latex_escape(analysis.get('summary') or analysis.get('results_detail', {}).get('summary'))
        dataset_name = cls._latex_escape(analysis.get('dataset_name') or analysis.get('dataset_id') or 'Dataset')
        generated_at = datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')
        probabilities = analysis.get('probabilities') or {}
        effect_sizes = analysis.get('effect_sizes') or analysis.get('results_detail', {}).get('effect_sizes') or {}
        plots = analysis.get('plots') or analysis.get('results_detail', {}).get('plots') or []

        body = r"""\documentclass{article}
\usepackage[margin=1in]{geometry}
\usepackage{longtable}
\usepackage{booktabs}
\usepackage{array}
\usepackage{float}
\begin{document}
\section*{StatMate Analysis Report}
"""
        body += f'\\textbf{{Dataset:}} {dataset_name}\\\\\n'
        body += f'\\textbf{{Generated:}} {generated_at}\\\\\n'
        body += '\\subsection*{Summary}\n'
        body += f"{summary or 'No summary available.'}\n"

        graph_assets = (
            (analysis.get('workflow_graph') or {}).get('assets')
            or analysis.get('results_detail', {}).get('workflow_graph', {}).get('assets')
            or {}
        )
        if graph_assets:
            body += '\\subsection*{Workflow graph}\n'
            body += '\\begin{figure}[H]\n\\centering\n'
            body += (
                '\\fbox{\\parbox{0.9\\linewidth}{\\centering Workflow graph preview '
                'is embedded in HTML/PDF exports.}}\\\\\n'
            )
            if graph_assets.get('alt'):
                body += f"\\textit{{{cls._latex_escape(graph_assets.get('alt'))}}}\n"
            body += '\\end{figure}\n'

        if probabilities:
            body += '\\subsection*{Statistical tests}\n'
            body += '\\begin{longtable}{p{0.35\\linewidth}p{0.25\\linewidth}p{0.25\\linewidth}}\n'
            body += '\\toprule\nTest & p-value & Callout \\\\\n\\midrule\n'
            for name, p_val in probabilities.items():
                escaped_name = cls._latex_escape(name)
                callout = 'Significant' if p_val < 0.05 else 'Not significant'
                body += f'{escaped_name} & {p_val:.4f} & {callout} \\\\\n'
            body += '\\bottomrule\n\\end{longtable}\n'

        if effect_sizes:
            body += '\\subsection*{Effect sizes}\n'
            body += '\\begin{longtable}{p{0.5\\linewidth}p{0.4\\linewidth}}\n'
            body += '\\toprule\nMetric & Value \\\\\n\\midrule\n'
            for name, val in effect_sizes.items():
                escaped_name = cls._latex_escape(name)
                body += f'{escaped_name} & {val:.3f} \\\\\n'
            body += '\\bottomrule\n\\end{longtable}\n'

        if plots:
            body += '\\subsection*{Plots}\n'
            for plot in plots:
                title = cls._latex_escape(plot.get('title') or 'Plot')
                description = cls._latex_escape(plot.get('description') or '')
                body += '\\begin{figure}[H]\n\\centering\n'
                body += (
                    f'\\fbox{{\\parbox{{0.9\\linewidth}}{{\\centering {title}'
                    f'\\\\[4pt]Images are included in the PDF/DOCX exports.}}}}\n'
                )
                if description:
                    body += f'\\caption*{{{description}}}\n'
                body += '\\end{figure}\n'

        body += '\\end{document}'
        return body.encode('utf-8')

    @staticmethod
    def render_docx(analysis: dict[str, Any]) -> bytes:
        """Render a small DOCX using python-docx."""
        try:
            from docx import Document
            from docx.shared import Inches
        except Exception as exc:  # pragma: no cover - optional dependency
            raise RuntimeError('python-docx is required for DOCX export') from exc

        doc = Document()
        doc.add_heading('StatMate Analysis Report', 0)
        doc.add_paragraph(f"Dataset: {analysis.get('dataset_name') or analysis.get('dataset_id')}")
        doc.add_paragraph(f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M UTC')}")

        summary = analysis.get('summary') or analysis.get('results_detail', {}).get('summary')
        doc.add_heading('Summary', level=1)
        doc.add_paragraph(summary or 'No summary available.')

        graph_assets = (
            (analysis.get('workflow_graph') or {}).get('assets')
            or analysis.get('results_detail', {}).get('workflow_graph', {}).get('assets')
            or {}
        )
        if graph_assets:
            doc.add_heading('Workflow graph', level=1)
            png_data = graph_assets.get('png_base64')
            svg_data = graph_assets.get('svg_base64')
            image_bytes = None
            if png_data:
                image_bytes = base64.b64decode(png_data)
            elif svg_data:
                try:
                    import cairosvg  # type: ignore

                    image_bytes = cairosvg.svg2png(bytestring=base64.b64decode(svg_data))
                except Exception:
                    image_bytes = None
            if image_bytes:
                doc.add_picture(io.BytesIO(image_bytes), width=Inches(5.5))
            else:
                doc.add_paragraph('Graph preview unavailable.')

        probabilities = analysis.get('probabilities') or {}
        if probabilities:
            doc.add_heading('Statistical tests', level=1)
            table = doc.add_table(rows=1, cols=3)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Test'
            hdr_cells[1].text = 'p-value'
            hdr_cells[2].text = 'Callout'
            for name, p_val in probabilities.items():
                row_cells = table.add_row().cells
                row_cells[0].text = name
                row_cells[1].text = f'{p_val:.4f}'
                row_cells[2].text = 'Significant' if p_val < 0.05 else 'Not significant'

        effect_sizes = analysis.get('effect_sizes') or analysis.get('results_detail', {}).get('effect_sizes') or {}
        if effect_sizes:
            doc.add_heading('Effect sizes', level=1)
            table = doc.add_table(rows=1, cols=2)
            hdr_cells = table.rows[0].cells
            hdr_cells[0].text = 'Metric'
            hdr_cells[1].text = 'Value'
            for name, val in effect_sizes.items():
                row_cells = table.add_row().cells
                row_cells[0].text = name
                row_cells[1].text = f'{val:.3f}'

        plots = analysis.get('plots') or analysis.get('results_detail', {}).get('plots') or []
        if plots:
            doc.add_heading('Plots', level=1)
            for plot in plots:
                title = plot.get('title') or 'Plot'
                description = plot.get('description') or ''
                doc.add_heading(title, level=2)
                image_b64 = plot.get('image_base64')
                if image_b64:
                    try:
                        image_data = base64.b64decode(image_b64)
                        doc.add_picture(io.BytesIO(image_data), width=Inches(6))
                    except Exception:
                        doc.add_paragraph('Image could not be rendered.')
                if description:
                    doc.add_paragraph(description)

        decision_steps = (
            analysis.get('decision_steps') or analysis.get('results_detail', {}).get('decision_steps') or []
        )
        doc.add_heading('Execution steps', level=1)
        if decision_steps:
            table = doc.add_table(rows=1, cols=5)
            hdr = table.rows[0].cells
            hdr[0].text = 'Step'
            hdr[1].text = 'Detail'
            hdr[2].text = 'Timestamp'
            hdr[3].text = 'p-value'
            hdr[4].text = 'Progress %'
            for step in decision_steps:
                row = table.add_row().cells
                row[0].text = str(step.get('step') or '')
                row[1].text = str(step.get('detail') or '')
                row[2].text = str(step.get('timestamp') or '')
                row[3].text = '' if step.get('p_value') is None else f"{step.get('p_value'):.4g}"
                row[4].text = str(step.get('progress_pct') or '')
        else:
            doc.add_paragraph('No steps captured.')

        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream.read()

    @staticmethod
    def render_csv(analysis: dict[str, Any]) -> bytes:
        """Render a compact CSV with probabilities and effect sizes."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(['metric', 'p_value', 'effect_size', 'callout'])
        probabilities = analysis.get('probabilities') or {}
        effect_sizes = analysis.get('effect_sizes') or analysis.get('results_detail', {}).get('effect_sizes') or {}
        keys = set(probabilities.keys()) | set(effect_sizes.keys())
        for name in keys:
            p_val = probabilities.get(name)
            effect_val = effect_sizes.get(name)
            callout = ''
            if p_val is not None:
                callout = 'Significant' if p_val < 0.05 else 'Not significant'
            writer.writerow(
                [name, p_val if p_val is not None else '', effect_val if effect_val is not None else '', callout]
            )
        return output.getvalue().encode('utf-8')

    @classmethod
    def build_repro_bundle(
        cls,
        analysis: dict[str, Any],
        dataset: pd.DataFrame | None = None,
        log_content: str | None = None,
    ) -> bytes:
        """Package data, logs, and decision context into a reproducibility bundle."""
        buffer = io.BytesIO()
        detail = analysis.get('results_detail') or {}
        decision_steps = analysis.get('decision_steps') or detail.get('decision_steps') or []
        assumption_log = analysis.get('assumption_log') or detail.get('assumption_log') or []
        test_hierarchy = analysis.get('test_hierarchy') or detail.get('test_hierarchy')
        workflow_graph = analysis.get('workflow_graph') or detail.get('workflow_graph') or {}

        snippets = []
        for name, p_val in (analysis.get('probabilities') or {}).items():
            snippets.append(
                f'# {name}\n'
                f'# Observed p-value: {p_val:.4f}\n# Replace <data> with your arrays\n# Example using scipy:\n'
                f'from scipy import stats\n# result = stats.ttest_ind(<group_a>, <group_b>, equal_var=False)\n'
            )

        with zipfile.ZipFile(buffer, 'w', compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('summary.md', analysis.get('summary') or detail.get('summary') or 'No summary provided.')
            zf.writestr('decision_steps.json', json.dumps(decision_steps, indent=2, default=str))
            zf.writestr('assumption_log.json', json.dumps(assumption_log, indent=2, default=str))
            if test_hierarchy:
                zf.writestr('test_hierarchy.json', json.dumps(test_hierarchy, indent=2, default=str))
            graph_assets = (workflow_graph or {}).get('assets') or {}
            if graph_assets.get('svg_base64'):
                zf.writestr('workflow_graph.svg', base64.b64decode(graph_assets['svg_base64']))
            elif graph_assets.get('svg'):
                zf.writestr('workflow_graph.svg', graph_assets['svg'])
            if graph_assets.get('png_base64'):
                zf.writestr('workflow_graph.png', base64.b64decode(graph_assets['png_base64']))
            if log_content:
                zf.writestr('execution.log', log_content)
            if dataset is not None:
                sample = dataset.head(200)
                masked_sample, mask_report = mask_dataframe(sample)
                zf.writestr('data_sample.csv', masked_sample.to_csv(index=False))
                if mask_report.get('masked_columns'):
                    zf.writestr('mask_report.json', json.dumps(mask_report, indent=2))
            zf.writestr('code_snippets.md', '\n'.join(snippets) if snippets else 'No statistical calls recorded.')

        buffer.seek(0)
        return buffer.getvalue()
